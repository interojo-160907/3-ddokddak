"""Exercise filters, BOM traversal, and Excel exports using isolated actual data."""
import os
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["DDOKDDAK_PROD3_PREVIEW"] = "1"
os.environ["DDOKDDAK_PROD3_PREVIEW_DATA_DIR"] = str(ROOT / "validation-data")
import json
from PySide6.QtWidgets import QApplication
from services.process_status_service import ProcessStatusService
from services.live_production_need_service import LiveProductionNeedService
from services.process_excel_exporter import export_overview_workbook, export_process_workbook, export_lot_work_order_workbook, export_hydration_lot_work_order_workbook
from services.inventory_status_service import InventoryStatusService, export_inventory
from services.lot_work_order_service import LotWorkOrderService
from services.bom_explorer import BomExplorerService
from ui.process_overview_page import ProcessOverviewPage
from ui.lot_work_order_page import INJECTION_COLUMNS, LOT_COLUMNS, SAME_CODE_COLUMNS, HYDRATION_SINGLE_COLUMNS, HYDRATION_SPLIT_COLUMNS
from openpyxl import load_workbook

app = QApplication([])
out = ROOT / "validation-artifacts" / "excel"
out.mkdir(parents=True, exist_ok=True)
report = {"filters": [], "exports": [], "bom": {}}
for label, service in (("backward", ProcessStatusService()), ("live", LiveProductionNeedService())):
    source = service.load_rows()
    for process in (None, "사출", "분리", "하이드레이션", "접착", "누수규격"):
        page = ProcessOverviewPage(fixed_process=process, initial_rows=ProcessOverviewPage._normalize(source), monitor_changes=False, include_packaging=False)
        page.search.setText("GG0270")
        page._apply_market_view()
        selected = page._export_master_rows()
        assert selected and all(r["이니셜"] == "GG0270" for r in selected)
        page.reset_all_filters()
        assert len(page._export_master_rows()) == len(source)
        safety = next(b for b in page.market_buttons if b.property("market") == "안전")
        safety.click()
        selected = page._export_master_rows()
        assert selected and all(r["진행현황"] == "안전" for r in selected)
        page.reset_all_filters()
        for check in page.detail_page.code_checks.values():
            check.setChecked(True)
        page.detail_page.order_quantity_check.setChecked(True)
        app.processEvents()
        report["filters"].append({"source":label, "process":process, "search_reset_market_codes_order":True})
        master = page._export_master_rows()
        target = out / (label + "_" + (process or "overview") + ".xlsx")
        if process:
            detail, compact = page.detail_page.export_rows(master)
            saved = export_process_workbook(process, detail, compact, output_path=target)
        else:
            detail, order, product = page.detail_page.export_overview_rows(master)
            saved = export_overview_workbook(detail, order, product, output_path=target)
        report["exports"].append(str(saved))
        page.close()
        page.deleteLater()

lot = LotWorkOrderService()
for process, columns in (("사출",INJECTION_COLUMNS),("분리",LOT_COLUMNS),("접착",SAME_CODE_COLUMNS),("누수규격",SAME_CODE_COLUMNS)):
    rows = lot.load_rows(process)
    assert rows
    report["exports"].append(str(export_lot_work_order_workbook(process,rows,columns,output_path=out/("lot_"+process+".xlsx"))))
single = lot.load_rows("하이드레이션",allocation_mode="whole")
split = lot.load_rows("하이드레이션",allocation_mode="split")
report["exports"].append(str(export_hydration_lot_work_order_workbook(single,split,HYDRATION_SINGLE_COLUMNS,HYDRATION_SPLIT_COLUMNS,output_path=out/"lot_hydration.xlsx")))
inventory = InventoryStatusService().build({})
export_inventory(inventory,out/"inventory.xlsx")
report["exports"].append(str(out/"inventory.xlsx"))
for path in report["exports"]:
    workbook = load_workbook(path,read_only=True,data_only=True)
    assert workbook.sheetnames and all(sheet.max_row > 1 for sheet in workbook)
    workbook.close()
bom = BomExplorerService()
for code in ("T1296","P1296","P3015"):
    candidates = bom.search(code)
    assert candidates, code
    graph = bom.graph(code)
    report["bom"][code] = {"search_matches":len(candidates),"graph_fields":list(graph)}
report["inventory_search"] = {}
for query in ("GG0270", "P1296", "Sincere", "GG0270, P1296"):
    result = InventoryStatusService().build({"search":query})
    assert result["rows"] > 0, query
    report["inventory_search"][query] = {"rows":result["rows"],"products":result["products"]}
(out.parent/"feature-verification.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(report,ensure_ascii=False,indent=2))
