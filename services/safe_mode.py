"""Isolated, immutable APS Excel snapshots; collectors always keep their normal paths."""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import time
from contextlib import closing
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

from config import DATA_CENTER_DIR, APP_VERSION
from services.program_gate import ProgramGate

SAFE_ROOT = Path(os.getenv("DDOKDDAK_SAFE_ROOT", str(DATA_CENTER_DIR / "안전모드")))
STATE = SAFE_ROOT / "적용정보.json"
_active: dict = {}
_restored = False
NAME_RE = re.compile(r"^(\d{6})_(오전|오후)\.xlsx$")
SNAPSHOT_REVISION = 3


def matching_spec(item: str, classification: str, power_match) -> str:
    """Optical identity within an exact sales name; process color aliases may differ."""
    suffix = item[power_match.start():]
    rest = item[power_match.end():]
    category = classification.casefold()
    if 'toric' in category:
        match = re.fullmatch(r'([+-]\d+\.\d{2})(\d{3})(?:[A-Za-z][A-Za-z0-9_]*)?', rest)
        if match:
            return power_match[0] + match[1] + match[2]
    elif 'm/f' in category or 'multi' in category:
        match = re.fullmatch(r'([+-]\d+\.\d{2})(?:[A-Za-z][A-Za-z0-9_]*)?', rest)
        if match:
            return power_match[0] + match[1]
    elif 'sph' in category and (not rest or re.fullmatch(r'[A-Za-z][A-Za-z0-9_]*', rest)):
        return power_match[0]
    return suffix


def file_key(name: str) -> tuple[str, int]:
    match = NAME_RE.fullmatch(name)
    if not match:
        raise ValueError("파일명은 YYMMDD_오전.xlsx 또는 YYMMDD_오후.xlsx여야 합니다.")
    day = datetime.strptime("20" + match[1], "%Y%m%d").date().isoformat()
    return day, int(match[2] == "오후")


def active() -> dict:
    global _restored, _active
    if not _restored:
        _restored = True
        try:
            candidate = json.loads(STATE.read_text(encoding="utf-8"))
            folder = SAFE_ROOT / candidate["digest"]
            if re.fullmatch(r"[0-9a-f]{64}", candidate["digest"]) and (folder / "aps.sqlite").is_file():
                _active = candidate
        except (OSError, ValueError, KeyError, TypeError):
            pass
    return dict(_active)


def database(default: Path) -> Path:
    info = active()
    return SAFE_ROOT / info["digest"] / "aps.sqlite" if info else default


def status_path(default: Path) -> Path:
    info = active()
    return SAFE_ROOT / info["digest"] / "스냅샷.json" if info else default


def _write(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def transient_control_error(exc: Exception) -> bool:
    if isinstance(exc, (requests.Timeout, requests.ConnectionError)):
        return True
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        return exc.response.status_code in {408, 429, 500, 502, 503, 504}
    return False


def fetch_control() -> dict:
    gate = ProgramGate(APP_VERSION)
    if not gate.endpoint:
        raise RuntimeError("전체설정 관리 API가 등록되지 않았습니다.")
    started = time.monotonic()
    diagnostic = {"checked_at": datetime.now().astimezone().isoformat()}
    try:
        response = requests.post(gate.endpoint, json={**gate.identity(), "action": "mode"}, timeout=(5, 30))
        response.raise_for_status()
        body = response.json()
        data = body.get("result", body)
        if body.get("ok") is False or data.get("mode") not in {"자동모드", "안전모드"}:
            raise RuntimeError(data.get("message") or "관리 API의 전체설정 모드 기능 업데이트가 필요합니다.")
        diagnostic.update(status="success", mode=data["mode"])
        return data
    except Exception as exc:
        # Do not persist URLs, PC identity, or response payloads.
        diagnostic.update(status="error", error_type=type(exc).__name__)
        raise
    finally:
        diagnostic["elapsed_seconds"] = round(time.monotonic() - started, 3)
        try:
            _write(DATA_CENTER_DIR / "settings" / "mode_check_status.json", diagnostic)
        except OSError:
            pass



def prepare(asset: dict) -> dict:
    """Build before publishing. No automatic database is modified."""
    name = str(asset.get("name") or "")
    day, half = file_key(name)
    if asset.get("local_path"):
        # Explicit development-only fixture; never read a local override in a packaged app.
        import sys
        if getattr(sys, "frozen", False) or os.getenv("DDOKDDAK_PROD3_PREVIEW") != "1":
            raise ValueError("로컬 안전모드 파일은 개발 미리보기에서만 사용할 수 있습니다.")
        content = Path(asset["local_path"]).read_bytes()
    else:
        url = str(asset.get("url") or "")
        if not re.fullmatch(r"https://raw\.githubusercontent\.com/interojo-160907/3-ddokddak/[0-9a-f]{40}/.+", url):
            raise ValueError("안전모드 파일의 고정 Git 주소가 올바르지 않습니다.")
        response = requests.get(url, timeout=(5, 45))
        response.raise_for_status()
        content = response.content
    if not content or len(content) > 50 * 1024 * 1024:
        raise ValueError("안전모드 파일 크기가 올바르지 않습니다.")
    expected = asset.get("git_sha")
    blob = hashlib.sha1(b"blob " + str(len(content)).encode() + b"\0" + content).hexdigest()
    if expected and expected != blob:
        raise ValueError("Git 파일 무결성 확인에 실패했습니다.")
    digest = hashlib.sha256(str(SNAPSHOT_REVISION).encode() + b'\0' + name.encode('utf-8') + b'\0' + content).hexdigest()
    folder = SAFE_ROOT / digest
    folder.mkdir(parents=True, exist_ok=True)
    source = folder / name
    source.write_bytes(content)
    info = {"digest": digest, "name": name, "label": f"{day} {'오후' if half else '오전'}", "selection_id": asset.get("selection_id", name), "snapshot_revision": SNAPSHOT_REVISION}
    if snapshot_valid(folder, info):
        return info
    build_snapshot(source, folder / "aps.building.sqlite", info)
    (folder / "aps.building.sqlite").replace(folder / "aps.sqlite")
    return info


def snapshot_valid(folder: Path, info: dict) -> bool:
    """Never reuse an incomplete or damaged cached conversion after an update."""
    try:
        status = json.loads((folder / "스냅샷.json").read_text(encoding="utf-8"))
        if status.get("digest") != info["digest"] or status.get("snapshot_revision") != SNAPSHOT_REVISION:
            return False
        with closing(sqlite3.connect((folder / "aps.sqlite").as_uri() + "?mode=ro", uri=True)) as con:
            return (con.execute("PRAGMA quick_check").fetchone()[0] == "ok"
                and con.execute("SELECT COUNT(*) FROM aps_plan").fetchone()[0] == status["stored_rows"]
                and status["stored_rows"] > 0)
    except (OSError, ValueError, KeyError, sqlite3.Error):
        return False


def build_snapshot(source: Path, target: Path, info: dict) -> None:
    from collectors.process_status_collector import _initialize, FIELDS
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(rows)]
        needed = ["설비 사이트 코드", "고객 이름", "이니셜", "수주번호", "신규분류 요약코드", "수요 제품 이름", "제품 코드", "납기일"]
        indexes = {name: headers.index(name) for name in needed}
        stages = {re.match(r"\[(\d+)\]", h)[1]: i for i, h in enumerate(headers) if re.match(r"\[(\d+)\]", h)}
        if not {"10", "20", "45", "55", "80"} <= stages.keys():
            raise ValueError("5개 필수 공정 열이 없습니다.")
        records = []
        counts = 0
        totals = {code: 0 for code in stages}
        for row_no, values in enumerate(rows, 2):
            row = {k: values[i] for k, i in indexes.items()}
            if str(row["설비 사이트 코드"] or "").strip() != "S관(3공장)":
                continue
            counts += 1
            item = str(row["제품 코드"] or "").strip()
            order = str(row["수주번호"] or "").strip()
            if not item or not order or not isinstance(row["납기일"], datetime):
                raise ValueError(f"{row_no}행의 제품코드·수주번호·납기일을 확인해 주세요.")
            initial = str(row["이니셜"] or "").strip()
            if "생산 수량" in headers:
                total = values[headers.index("생산 수량")]
                stage_values = [values[index] for index in stages.values()]
                if any(v not in (None, "") and (not isinstance(v, (int, float)) or not math.isfinite(v) or v < 0) for v in stage_values):
                    raise ValueError(f"{row_no}행의 공정 수량을 확인해 주세요.")
                if not isinstance(total, (int, float)) or not math.isfinite(total) or abs(sum(v or 0 for v in stage_values) - total) > 0.01:
                    raise ValueError(f"{row_no}행의 생산 수량 합계가 공정 합계와 다릅니다.")
            # APS exports its demand class in Initial: 국내, PB, 안전(...), or customer initial.
            demand_type = "안전" if "안전" in initial else initial if initial in {"국내", "PB"} else "이니셜"
            power_match = re.search(r"[+-]\d{2}\.\d{2}", item)
            # R/Q/P are process-specific codes, not separate demand products.
            # Match only within the same source demand/name/class/date and exact
            # optical specification (CP, AXIS, ADD); color aliases are scoped by sales name.
            demand_key = item
            if power_match and item[:1] in {'R', 'Q', 'P'}:
                suffix = matching_spec(item, str(row['신규분류 요약코드'] or ''), power_match)
                identity = [order, initial, str(row['고객 이름'] or ''),
                    str(row['신규분류 요약코드'] or ''), str(row['수요 제품 이름'] or ''),
                    row['납기일'].date().isoformat(), suffix]
                digest_key = hashlib.sha256(json.dumps(identity, ensure_ascii=False).encode('utf-8')).hexdigest()[:24]
                demand_key = f'EXCEL:{digest_key}:{suffix}'
            for code, index in stages.items():
                value = values[index]
                if value in (None, ""):
                    continue
                if not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                    raise ValueError(f"{row_no}행 공정 {code} 수량 오류")
                totals[code] += value
                if not value or code not in {"10", "20", "45", "55", "80", "85"}:
                    continue
                record = {"res_site_id": "S관(3공장)", "so_id": order, "initial": initial,
                    "cust_name": str(row["고객 이름"] or ""), "demand_group_id": str(row["신규분류 요약코드"] or ""),
                    "demand_item_id": demand_key, "item_id": item, "item_cd": demand_key,
                    "item_name": str(row["수요 제품 이름"] or ""),
                    "demand_item_name": str(row["수요 제품 이름"] or ""),
                    "due_date": row["납기일"].date().isoformat(), "oper_id": code, "plan_qty": value,
                    "power": power_match[0] if power_match else "", "demand_type": demand_type,
                    "seq": row_no, "target_datetime": info["label"]}
                records.append((*[record.get(field) for field in FIELDS], json.dumps(record, ensure_ascii=False)))
        if not records:
            raise ValueError("S관 공정 데이터가 없습니다.")
        target.unlink(missing_ok=True)
        with closing(sqlite3.connect(target)) as con:
            _initialize(con)
            con.executemany(f"INSERT INTO aps_plan ({','.join(FIELDS)},payload_json) VALUES ({','.join('?' for _ in range(len(FIELDS)+1))})", records)
            con.execute("INSERT INTO sync_meta VALUES (1,?,?,?,?,?,?,?)", ("S관", counts, len(records), len({r[FIELDS.index('so_id')] for r in records}), info["label"], datetime.now().isoformat(), str(source)))
            con.commit()
            if con.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise ValueError("스냅샷 DB 검증 실패")
        _write(target.parent / "스냅샷.json", {**info, "status": "success", "source_refreshed_at": info["label"], "stored_rows": len(records), "source_rows": counts, "process_totals": totals, "source": "안전모드 Excel"})
    finally:
        book.close()


def activate(info: dict) -> None:
    global _active, _restored
    _write(STATE, info)
    _active, _restored = dict(info), True


def deactivate() -> None:
    global _active, _restored
    STATE.unlink(missing_ok=True)
    _active, _restored = {}, True


def cleanup() -> None:
    """Remove only owned hash directories after the UI has switched away."""
    if active() or not SAFE_ROOT.exists():
        return
    root = SAFE_ROOT.resolve()
    for folder in root.iterdir():
        if folder.is_dir() and not folder.is_symlink() and re.fullmatch(r"[0-9a-f]{64}", folder.name) and folder.resolve().parent == root:
            shutil.rmtree(folder)
