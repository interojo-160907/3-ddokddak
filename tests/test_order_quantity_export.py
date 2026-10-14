import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from services.process_excel_exporter import (
    PROCESS_ORDER, build_process_export_payload, build_overview_export_payload,
    export_overview_workbook, export_process_workbook,
)
from ui.process_overview_page import ProcessOverviewPage
from ui.live_production_need_page import LiveProductionNeedPage


def row(order="O1", quantity=1234):
    return {
        "수주번호": order, "이니셜": "A", "품목코드": "T1", "T코드": "T1",
        "P코드": "P1", "Q코드": "Q1", "R코드": "R1",
        "품명판매": "판매품명", "품명P": "생산품명", "품명Q": "분리품명", "품명R": "사출품명",
        "POWER": "-01.00", "납기일": "2026-11-01", "신규분류요약": "1-Day_Sph",
        "오더수량PCS": quantity, "공정": {name: 100 for name in PROCESS_ORDER},
    }


class OrderQuantityExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_every_page_and_toggle_exports_numeric_quantity_and_group_sum(self):
        sources = [row(), row("O2", 2000)]
        for page_class in (ProcessOverviewPage, LiveProductionNeedPage):
            for process in (None, *PROCESS_ORDER):
                with self.subTest(page=page_class.__name__, process=process):
                    kwargs = {"fixed_process": process, "initial_rows": sources}
                    if page_class is ProcessOverviewPage:
                        kwargs["monitor_changes"] = False
                    page = page_class(**kwargs)
                    try:
                        snapshots = []
                        for checked in (False, True):
                            page.detail_page.order_quantity_check.setChecked(checked)
                            if process:
                                details, compact = page.detail_page.export_rows(sources)
                                payload = build_process_export_payload(process, details, compact)
                            else:
                                details, orders, products = page.detail_page.export_overview_rows(sources)
                                payload = build_overview_export_payload(details, orders, products)
                            snapshots.append(payload)
                            for sheet in payload["sheets"]:
                                headers = sheet["columns"]
                                index = headers.index("오더수량")
                                self.assertNotIn("오더수량", sheet["hiddenColumns"])
                                previous = process if sheet["name"] == "납기별 상세" and process else "누수규격"
                                if process and sheet["name"] == "간략히보기":
                                    previous = f"{process} 부족수량"
                                self.assertEqual(headers[index - 1], previous)
                                values = [data[index] for data in sheet["rows"]]
                                self.assertEqual(values, [3234] if process and sheet["name"] == "간략히보기" else [1234, 2000])
                        self.assertEqual(snapshots[0], snapshots[1])
                    finally:
                        page.close()
                        page.deleteLater()

    def test_workbooks_preserve_numeric_zero_and_unknown_blank(self):
        rows = [row(quantity=1234), row("O2", 0), row("O3", None)]
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "overview.xlsx"
            export_overview_workbook(rows, rows, rows, output_path=output)
            self._assert_workbook(output)
            for process in PROCESS_ORDER:
                details = [{**source, **source["공정"]} for source in rows]
                output = Path(folder) / f"{process}.xlsx"
                export_process_workbook(process, details, rows, output_path=output)
                self._assert_workbook(output)

    def _assert_workbook(self, output):
        workbook = load_workbook(output)
        try:
            for sheet in workbook:
                headers = [cell.value for cell in sheet[4]]
                column = headers.index("오더수량") + 1
                cells = [sheet.cell(i, column) for i in (5, 6, 7)]
                self.assertEqual([cell.value for cell in cells], [1234, 0, None])
                self.assertEqual([cell.data_type for cell in cells[:2]], ["n", "n"])
                self.assertEqual(cells[0].number_format, "#,##0")
        finally:
            workbook.close()

    def test_grouping_deduplicates_and_preserves_unknown(self):
        page = ProcessOverviewPage(initial_rows=[], monitor_changes=False)
        try:
            detail = page.detail_page
            source = row()
            grouped = detail._group_main_rows([source, source], "order", "판매")
            self.assertEqual(grouped[0]["오더수량PCS"], 1234)
            detail.fixed_process = "사출"
            grouped = detail._group_same_process_code([source, row("O2", None)], "R")
            self.assertIsNone(grouped[0]["오더수량PCS"])
        finally:
            page.close()
            page.deleteLater()
